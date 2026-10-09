"""Synthetic Chinese Office fixtures; no user data or network access."""
from pathlib import Path
import sys
from docx import Document
from docx.shared import Inches as DocInches
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from PIL import Image, ImageDraw

root = Path(sys.argv[1]); root.mkdir(parents=True, exist_ok=True)
image = Image.new("RGB", (800, 260), "#152638")
draw = ImageDraw.Draw(image)
for index, height in enumerate([100, 150, 205]):
    draw.rectangle((90 + index * 180, 235-height, 190 + index*180, 235), fill="#56B9EC")
image.save(root / "chart.png")
d = Document()
style = d.styles["Normal"]; style.font.name = "Microsoft YaHei"
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
d.add_heading("星辰文档验收 / Office acceptance", 0)
d.add_paragraph("中文排版测试：表格、图片与段落应保持可读。HARNESS_DOCUMENT_MARKER")
table = d.add_table(rows=3, cols=3); table.style = "Table Grid"
for row, values in zip(table.rows, [("项目", "数量", "金额"), ("文档", "2", "125.50"), ("演示", "3", "240.00")]):
    for cell, value in zip(row.cells, values): cell.text = value
d.add_picture(str(root / "chart.png"), width=DocInches(5.5))
d.save(root / "中文 文档.docx")
r = Presentation(); r.slide_width = Inches(13.333); r.slide_height = Inches(7.5)
s = r.slides.add_slide(r.slide_layouts[0]); s.shapes.title.text = "星辰演示 / Slides"
s.placeholders[1].text = "HARNESS_PRESENTATION_MARKER"
s = r.slides.add_slide(r.slide_layouts[5]); s.shapes.title.text = "季度数据"
data = CategoryChartData(); data.categories = ["第一季度", "第二季度", "第三季度"]; data.add_series("销售额", (10, 20, 35))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(10), Inches(4.5), data)
s = r.slides.add_slide(r.slide_layouts[5]); s.shapes.title.text = "图片与文字"
s.shapes.add_picture(str(root / "chart.png"), Inches(2), Inches(2), width=Inches(8))
for slide in r.slides:
    for shape in slide.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs: run.font.name = "Microsoft YaHei"; run.font.size = Pt(24)
r.save(root / "中文 演示.pptx")
w = Workbook(); sheet = w.active; sheet.title = "销售"
sheet.append(["数量", "单价", "收入"]); sheet.append([2, 12.5, "=A2*B2"]); sheet.append([3, 20, "=A3*B3"])
sheet["C4"] = "=SUM(C2:C3)"; sheet["D2"] = '=IF(C4=85,"通过","失败")'
chart = BarChart(); chart.add_data(Reference(sheet, min_col=3, min_row=1, max_row=3), titles_from_data=True); sheet.add_chart(chart, "F2")
w.save(root / "中文 表格.xlsx")
sheet["E2"] = "=1/0"; w.save(root / "公式错误.xlsx")
print("Synthetic fixtures ready")
